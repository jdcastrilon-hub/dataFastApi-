import io
from datetime import datetime
from io import BytesIO
from typing import Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import desc, text
from sqlalchemy.orm import Session, joinedload

from . import models
from app.core.numeradores import repository_numerador
from app.modules.stock.articulos import model_articulos
from app.modules.comercial.listaprecio import model_listaprecio

MAX_LOGS_AUDITORIA = 10
CODIGO_NUMERADOR_CARGAPRECIOS = "CARGAPRECIOS"

COLUMNAS_OBLIGATORIAS = ["codigo_articulo", "precio_venta"]

DESCRIPCION_COLUMNAS = {
    "codigo_articulo": ("Si", "Codigo del articulo (el mismo que ya usa en el resto del sistema) - debe existir previamente, esta carga NO crea articulos nuevos."),
    "precio_venta": ("Si", "Precio de venta a cargar contra la lista elegida. Debe ser mayor o igual a 0."),
}

RELLENO_OBLIGATORIA = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")


def generar_excel_plantilla() -> BytesIO:
    """Plantilla en blanco con las 2 columnas que espera la carga de precios,
    mas una pestaña de instrucciones (mismo shape que cargastock)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Plantilla"

    encabezados = COLUMNAS_OBLIGATORIAS
    ws.append(encabezados)
    for celda in ws[1]:
        celda.font = Font(bold=True)
        celda.fill = RELLENO_OBLIGATORIA

    for columna in ws.columns:
        longitud = max(len(str(celda.value)) if celda.value is not None else 0 for celda in columna)
        ws.column_dimensions[columna[0].column_letter].width = min(max(longitud + 4, 14), 40)

    ws.freeze_panes = "A2"

    ws_info = wb.create_sheet("Instrucciones")
    ws_info.append(["Columna", "Obligatoria", "Detalle"])
    for celda in ws_info[1]:
        celda.font = Font(bold=True)

    for columna in encabezados:
        obligatoria, detalle = DESCRIPCION_COLUMNAS[columna]
        fila = ws_info.max_row + 1
        ws_info.append([columna, obligatoria, detalle])
        ws_info.cell(row=fila, column=1).fill = RELLENO_OBLIGATORIA

    ws_info.column_dimensions["A"].width = 18
    ws_info.column_dimensions["B"].width = 14
    ws_info.column_dimensions["C"].width = 90
    for fila_info in ws_info.iter_rows(min_row=2):
        for celda in fila_info:
            celda.alignment = celda.alignment.copy(wrap_text=True, vertical="top")

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def _limitar_logs(logs):
    if not logs:
        return logs
    return logs[-MAX_LOGS_AUDITORIA:]


def _leer_excel(contenido: bytes):
    libro = load_workbook(io.BytesIO(contenido), data_only=True)
    hoja = libro.active
    filas = list(hoja.iter_rows(values_only=True))
    if not filas:
        return [], []
    encabezado = [str(c).strip().lower() if c is not None else "" for c in filas[0]]
    return encabezado, filas[1:]


def _texto(valor) -> str:
    if valor is None:
        return ""
    return str(valor).strip()


def _numero(valor):
    try:
        if valor is None or valor == "":
            return None
        return float(valor)
    except (TypeError, ValueError):
        return None


def procesar_carga_precios(
    db: Session,
    id_lista: int,
    fecha_carga,
    observacion: Optional[str],
    nombre_archivo: str,
    contenido: bytes,
    confirmar: bool,
    usuario_nombre: str,
    id_emp: int,
):
    """Fase 1: parsea y valida el archivo completo contra la lista elegida - cada
    codigo_articulo debe existir en m_articulos de la misma empresa (a diferencia
    de cargastock, esta carga NUNCA crea articulos nuevos). Fase 2 (solo si
    confirmar=True y la fase 1 quedo limpia): inserta cabezal/detalle e impacta
    p_precios con un INSERT directo - el trigger ya existente (ver alembic
    9c8225580f88) actualiza el snapshot y cascada las reglas de categoria solo,
    sin necesitar un stored procedure propio (mismo criterio ya usado en
    repository_ajusteprecio.py)."""
    lista = db.query(model_listaprecio.MListaPrecio).filter(
        model_listaprecio.MListaPrecio.id_lista == id_lista,
        model_listaprecio.MListaPrecio.id_emp == id_emp,
        model_listaprecio.MListaPrecio.activo.is_(True)
    ).first()
    if lista is None:
        return {"status": "error", "message": "La lista de precios seleccionada no existe o no esta activa.", "errores": [], "resumen": None}

    encabezado, filas_excel = _leer_excel(contenido)
    if not filas_excel:
        return {"status": "error", "message": "El archivo esta vacio o no tiene filas de datos.", "errores": [], "resumen": None}

    faltantes = [c for c in COLUMNAS_OBLIGATORIAS if c not in encabezado]
    if faltantes:
        return {
            "status": "error",
            "message": "El archivo no tiene las columnas esperadas.",
            "errores": [{"fila": 1, "mensaje": f"Falta la columna '{c}'"} for c in faltantes],
            "resumen": None
        }

    idx = {nombre: pos for pos, nombre in enumerate(encabezado) if nombre}

    # Precarga de articulos de la empresa (una sola consulta, no por fila).
    articulos = db.query(model_articulos.Articulo).filter(model_articulos.Articulo.id_emp == id_emp).all()
    articulo_por_codigo = {}
    codigos_ambiguos = set()
    for a in articulos:
        if not a.cod_articulo:
            continue
        clave = a.cod_articulo.strip().lower()
        if clave in articulo_por_codigo and articulo_por_codigo[clave] != a.id_articulo:
            codigos_ambiguos.add(clave)
        articulo_por_codigo[clave] = a.id_articulo

    errores = []
    filas_procesadas = []
    codigos_vistos = {}  # codigo (normalizado) -> fila donde aparecio primero

    for i, fila in enumerate(filas_excel, start=2):  # la fila 1 del excel es el encabezado
        if fila is None or all(c is None for c in fila):
            continue

        def val(nombre_col):
            pos = idx.get(nombre_col)
            return fila[pos] if pos is not None and pos < len(fila) else None

        codigo_txt = _texto(val("codigo_articulo"))
        precio_num = _numero(val("precio_venta"))

        if not codigo_txt:
            errores.append({"fila": i, "mensaje": "El codigo de articulo es obligatorio."})
            continue
        if precio_num is None or precio_num < 0:
            errores.append({"fila": i, "mensaje": "El precio de venta es obligatorio y debe ser mayor o igual a 0."})
            continue

        clave = codigo_txt.strip().lower()
        if clave in codigos_ambiguos:
            errores.append({"fila": i, "mensaje": f"El codigo '{codigo_txt}' esta asociado a mas de un articulo; corrijalo en el maestro de Articulos antes de continuar."})
            continue

        id_articulo = articulo_por_codigo.get(clave)
        if id_articulo is None:
            errores.append({"fila": i, "mensaje": f"El articulo con codigo '{codigo_txt}' no existe. Esta carga no crea articulos nuevos."})
            continue

        if clave in codigos_vistos:
            errores.append({"fila": i, "mensaje": f"El codigo '{codigo_txt}' ya aparece en la fila {codigos_vistos[clave]}; elimine la fila duplicada."})
            continue
        codigos_vistos[clave] = i

        filas_procesadas.append({"fila": i, "id_articulo": id_articulo, "precio_venta": precio_num})

    if errores:
        return {
            "status": "error",
            "message": "El archivo tiene errores, corrijalos y vuelva a intentarlo.",
            "errores": sorted(errores, key=lambda e: e["fila"]),
            "resumen": None
        }

    if not filas_procesadas:
        return {"status": "error", "message": "El archivo no tiene filas validas para procesar.", "errores": [], "resumen": None}

    resumen = {
        "totalFilas": len(filas_procesadas),
        "articulosActualizados": len({f["id_articulo"] for f in filas_procesadas}),
    }

    if not confirmar:
        # Solo vista previa: nada se toca en la base de datos todavia.
        return {"status": "success", "message": "Archivo validado correctamente.", "errores": [], "resumen": resumen}

    # --- Fase 2: confirmar y grabar (todo o nada) ---
    try:
        nro_docum = repository_numerador.siguiente_numerador(db, id_emp, CODIGO_NUMERADOR_CARGAPRECIOS)

        ids_articulo = [f["id_articulo"] for f in filas_procesadas]
        filas_precio_actual = db.execute(
            text("SELECT id_articulo, precio_venta FROM s_precioxarticulo WHERE id_lista = :lista AND id_articulo = ANY(:ids)"),
            {"lista": id_lista, "ids": ids_articulo}
        ).all()
        precios_actuales = {r.id_articulo: r.precio_venta for r in filas_precio_actual}

        bd_cabecera = models.CargaPrecios(
            id_emp=id_emp,
            id_lista=id_lista,
            documento="cargaprec",
            nro_docum=nro_docum,
            fecha_carga=fecha_carga,
            observacion=observacion,
            nombre_archivo=nombre_archivo,
            vista="CargaPrecios",
            fecha_mod=datetime.now(),
            logs=_limitar_logs([{
                "operacion": "Nuevo",
                "usuario_mod": usuario_nombre,
                "fecha_mod": str(datetime.now())
            }])
        )
        db.add(bd_cabecera)
        db.flush()

        for linea, f in enumerate(filas_procesadas, start=1):
            db.add(models.DetalleCargaPrecios(
                id_trans=bd_cabecera.id_trans,
                id_articulo=f["id_articulo"],
                linea=linea,
                precio_venta=f["precio_venta"]
            ))

            # Impacta p_precios: el trigger ins_p_precios actualiza el snapshot
            # s_precioxarticulo, y si la lista elegida es la General, ins_p_precios_reglas
            # cascada a las listas de cliente con regla activa (ver 9c8225580f88).
            db.execute(text("""
                INSERT INTO p_precios (id_trans, linea, id_emp, id_lista, id_articulo,
                    precio_anterior, precio_nuevo, origen, usuario_mod, fecha_mod)
                VALUES (:id_trans, :linea, :id_emp, :id_lista, :id_articulo,
                    :precio_anterior, :precio_nuevo, 'CARGA_PRECIOS', :usuario_mod, :fecha_mod)
            """), {
                "id_trans": bd_cabecera.id_trans,
                "linea": linea,
                "id_emp": id_emp,
                "id_lista": id_lista,
                "id_articulo": f["id_articulo"],
                "precio_anterior": precios_actuales.get(f["id_articulo"]),
                "precio_nuevo": f["precio_venta"],
                "usuario_mod": usuario_nombre,
                "fecha_mod": datetime.now()
            })

        db.commit()
        db.refresh(bd_cabecera)

        return {
            "status": "success",
            "message": "Carga de precios procesada exitosamente.",
            "errores": [],
            "resumen": resumen,
            "idTrans": bd_cabecera.id_trans
        }

    except Exception:
        db.rollback()
        raise


# Paginacion
def get_cargas_paginated(db: Session, page: int, size: int, id_emp: int, texto: str = None):
    query = db.query(models.CargaPrecios).filter(models.CargaPrecios.id_emp == id_emp)

    if texto:
        patron = f"%{texto}%"
        query = query.filter(
            (models.CargaPrecios.observacion.ilike(patron)) | (models.CargaPrecios.nombre_archivo.ilike(patron))
        )

    total_records = query.count()
    offset = page * size
    items = (
        query
        .options(joinedload(models.CargaPrecios.lista))
        .order_by(desc(models.CargaPrecios.fecha_mod))
        .offset(offset)
        .limit(size)
        .all()
    )
    total_pages = (total_records + size - 1) // size

    return {
        "content": items,
        "totalElements": total_records,
        "totalPages": total_pages,
        "number": page,
        "size": size
    }


# Obtener una carga por ID (para la pantalla de "ver")
def get_carga_precios(db: Session, id_trans: int):
    return (
        db.query(models.CargaPrecios)
        .filter(models.CargaPrecios.id_trans == id_trans)
        .options(
            joinedload(models.CargaPrecios.lista),
            joinedload(models.CargaPrecios.detalles).joinedload(models.DetalleCargaPrecios.articulo)
        )
        .first()
    )

# No hay funcion de eliminar a proposito (mismo criterio que cargastock): esta
# carga impacta p_precios y puede haber cascadeado reglas de categoria hacia
# otras listas - revertir con seguridad requeriria logica que este modulo no
# esta pensado para hacer.
